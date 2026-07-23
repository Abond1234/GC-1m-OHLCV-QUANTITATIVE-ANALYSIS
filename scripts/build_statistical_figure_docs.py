"""Build one quantitative interpretation DOCX for every statistical-research figure.

The source PNGs remain unchanged. Deliverable DOCX files are written to the
figure-pack root so that all interpretations are available in one directory.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
FIGURE_ROOT = PROJECT_ROOT / "reports" / "figures" / "statistical_feature_research"

NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(92, 101, 111)
LIGHT_GRAY = "F4F6F9"
RULE_GRAY = "D7DBE2"
POSITIVE = RGBColor(42, 157, 143)
CAUTION = RGBColor(122, 90, 0)
RISK = RGBColor(155, 28, 28)


@dataclass(frozen=True)
class FigureBrief:
    image: str
    title: str
    section: str
    metric_role: str
    takeaway: str
    represents: str
    reading: str
    interpretation: str
    implications: str
    limitations: str


BRIEFS: tuple[FigureBrief, ...] = (
    FigureBrief(
        "section02_data/01_trusted_source_coverage.png",
        "Trusted Source Coverage",
        "Section 2 — Trusted research data",
        "Data-integrity diagnostic; Sharpe ratio is not applicable",
        "GC and MGC provide nearly matched source depth and identical date coverage, giving the downstream study a balanced and auditable market-data foundation.",
        "The two panels summarize the raw trusted one-minute research table before any feature or label filtering. The left panel compares source-row counts in millions; the right compares New York trading-date coverage. GC contributes 1,759,671 rows and MGC 1,727,985 rows. Both span 1,556 New York trading dates from May 2021 through May 2026. This is a provenance and coverage check, not a statement about predictive value.",
        "Read the products side by side. Similar bar counts indicate comparable sampling density, while identical trading-date counts show that neither contract family has a shorter calendar history. Small row-count differences are expected because tradability, contract rolls, and recorded activity differ by product. A material date mismatch would have signalled a join, ingestion, or filtering problem before any quantitative research began.",
        "The main quantitative conclusion is that cross-product context can be studied without an obvious calendar-coverage confound. GC has roughly 1.8% more rows than MGC, but both histories share their reported endpoints and trading-date count. That makes later GC-only Phase 1 analysis a deliberate scope choice rather than a response to missing MGC history. The source table also retains rollover, liquidity, timezone, and tradability controls, so downstream observations can be filtered causally.",
        "Treat this figure as an admission gate. If coverage drifts on a rerun, stop before interpreting features: verify source version, exchange calendar handling, and product filters. For model governance, retain the row/date counts alongside the data hash or dataset version so a future notebook execution can be tied to the same research universe.",
        "Sharpe ratio would be meaningless here because there is no return stream, position rule, or time-series payoff. Large row counts also do not guarantee independence: one-minute bars are serially related and share trading days. The chart validates scope and continuity only; it does not validate prices, labels, execution assumptions, or model performance.",
    ),
    FigureBrief(
        "section03_population/02_eligible_population.png",
        "Eligible Population Composition",
        "Section 3 — Eligible observation frame",
        "Population diagnostic; Sharpe ratio is not applicable",
        "The research population is large but intentionally unbalanced by session and frozen into distinct Development, Validation, and Final-test partitions.",
        "This figure summarizes the 586,530 GC observations that survive the decision-to-entry eligibility contract. The left panel splits them into London (219,938) and New York (366,592) entries. The right panel shows the frozen partitions: 306,230 Development, 118,076 Validation, and 162,224 Final test. Every observation maps a valid decision bar to the true next one-minute entry bar under contract, segment, roll, liquidity, date, and session constraints.",
        "Bar height is a count, not a performance score. The New York session is larger because its entry window is five hours versus three hours for London. The partition panel should be read as research governance: Development is used for fitting and selection, Validation for honest confirmation, and Final test for the one-time governed read. The unequal bars are therefore expected and should not be mechanically equalized.",
        "The sample is sufficiently broad for date-level evaluation, but the unit of evidence is not 586,530 independent rows. Intraday observations within a trading date share market state and overlapping forward paths. That is why later confidence intervals and feature diagnostics aggregate or resample by New York trading date. Session-specific reporting is also essential: a pooled result would be weighted toward New York and could conceal different behaviour in London.",
        "Use the chart to police leakage and weighting. Any model comparison should preserve the frozen partitions and report London and New York separately before pooling. Final-test counts must never influence feature selection, thresholds, hyperparameters, or chart rankings. When presenting pooled statistics, state the weighting rule explicitly so the larger New York population does not silently determine the conclusion.",
        "Sharpe ratio is not applicable because these are eligible observations rather than realized sequential returns. Counts say nothing about economic edge. They also do not reveal effective sample size after date clustering, class imbalance, or label availability at each horizon; those questions are handled by later figures.",
    ),
    FigureBrief(
        "section04_labels/03_label_availability.png",
        "Forward-Label Availability",
        "Section 4 — Forward outcome labels",
        "Label-quality diagnostic; Sharpe ratio is not applicable",
        "Complete-path availability remains very high at every horizon, declining monotonically from 99.99% at 5 minutes to 98.74% at 180 minutes.",
        "The line plots the fraction of all 586,530 eligible observations for which the entire future path is valid at each governed horizon: 5, 15, 30, 60, 120, and 180 minutes. Availability is 99.989% at 5 minutes and 98.745% at 180 minutes, where 579,166 observations remain and 7,364 are unavailable. Paths are not shortened; any gap, contract or segment change, date change, tradability failure, forced-exit breach, or invalid price makes the label unavailable.",
        "Read the downward slope as mechanical path attrition. Longer labels require more consecutive valid bars and are more likely to encounter the forced-exit boundary or a data discontinuity. The y-axis is deliberately tight, so the visual decline looks larger than its absolute size. The annotations provide the correct scale: even the longest horizon retains almost 99% of observations.",
        "The high rates reduce the risk that model results are driven by broad missing-label selection. Nevertheless, the missingness is horizon- and time-dependent: late-session entries are more exposed to the 15:30 New York forced exit, and London/New York availability differs slightly. Therefore every horizon must use its own availability mask and trading-date count. Comparing metrics across horizons without acknowledging different complete-path populations can produce small but systematic selection effects.",
        "Use availability as a denominator audit in every downstream table. A sudden drop after code or data changes should trigger investigation before IC, AUC, or Sharpe is recalculated. Preserve the first-failure reason because it distinguishes normal horizon attrition from bad joins, missing minutes, roll leakage, or corrupted prices.",
        "Sharpe ratio is not a label-quality metric. It requires a signed return series from a trading policy, whereas this chart measures whether future outcomes can be computed. High availability does not imply profitable or unbiased labels, and the many overlapping forward paths are not independent observations.",
    ),
    FigureBrief(
        "section05_baseline/04_return_quantiles.png",
        "Unconditional Return Quantiles",
        "Section 5 — Baseline behaviour",
        "Descriptive event study; Sharpe ratio is not applicable",
        "Forward-return dispersion expands rapidly with horizon while the distribution centre remains close to zero, implying opportunity without an unconditional directional edge.",
        "The chart traces the 5th, 25th, median, 75th, and 95th percentiles of raw GC forward returns in ticks over six horizons. At 5 minutes the central 90% runs roughly from -31 to +31 ticks and the median is zero. By 60 minutes the 5th and 95th percentiles are approximately -111 and +108 ticks, while the mean is only -0.54 tick. At 180 minutes the standard deviation is about 140 ticks and the mean remains below one tick in magnitude.",
        "Read the outer lines as tail expansion and the inner lines as the interquartile range. Their separation increases with horizon, as expected when more price movement accumulates. The near-zero median and broadly symmetric positive and negative bands show that unconditional direction is weak. Slight asymmetries should be judged relative to the much larger dispersion, not by sign alone.",
        "The baseline establishes the hurdle for engineered features. A useful directional feature must condition this wide distribution so that the signed mean, rank relationship, or executable expectancy changes reliably out of sample. A useful expansion feature need not predict sign; it can forecast when the range will be large. This distinction becomes central later: volatility and clock features strongly predict expansion even though the standalone directional strategy fails.",
        "Use these quantiles for scale calibration, stop/target plausibility, robust outlier checks, and economic-materiality thresholds. They also warn against evaluating a model only with mean squared error: the target is heavy-tailed and heteroskedastic. Session- and time-conditioned baselines should be consulted before treating a feature as novel.",
        "These are overlapping forward outcomes measured at every eligible minute, not a sequence of non-overlapping portfolio returns. A Sharpe ratio would overstate effective sample size and misrepresent descriptive labels as a strategy. The chart also excludes transaction costs, position overlap, sizing, and execution, so it cannot support a trading claim.",
    ),
    FigureBrief(
        "section05_baseline/04_return_quantiles_and_uncertainty.png",
        "Return Quantiles and Date-Block Uncertainty",
        "Section 5 — Baseline behaviour",
        "Descriptive event study with date-block inference; Sharpe ratio is not applicable",
        "The distribution widens materially with horizon, but the daily-weighted mean remains economically small relative to its uncertainty and tail width.",
        "This two-panel figure combines distribution shape with inference. The left panel repeats raw forward-return quantiles across 5–180 minutes. The right panel plots the daily-weighted mean return at each horizon with a date-block bootstrap interval. Weighting dates rather than individual intraday rows prevents unusually active days from dominating and acknowledges within-day dependence.",
        "On the left, focus on the distance between p05 and p95 rather than the small movement of the median. On the right, compare each point and interval with the zero line. An interval overlapping or remaining close to zero means the apparent unconditional mean is not a robust directional benchmark, even when millions of minute-level paths contribute. The date block is the relevant uncertainty unit because paths overlap heavily within a session.",
        "The two panels explain why raw sample size can be deceptive. The tails expand from tens to well over one hundred ticks, while the location estimate is tiny. Consequently, a model can achieve a visually large range forecast or rank IC without creating a positive signed return after costs. Conversely, small mean differences should not be promoted unless their date-level interval, economic magnitude, and Validation behaviour all agree.",
        "Use this figure as the pre-feature null model. Later IC and spread results should be compared with the unconditional scale and assessed using the same date-level discipline. It also justifies robust statistics, date-block bootstrap intervals, and separate expansion versus direction targets.",
        "Sharpe is inappropriate because the plotted returns overlap and no trading rule converts them into a sequential daily P&L stream. The bootstrap interval quantifies uncertainty in a descriptive mean; it is not a strategy risk-adjusted return estimate. Transaction costs and capacity are absent.",
    ),
    FigureBrief(
        "section05_baseline/05_range_and_realized_volatility.png",
        "Range and Realized Volatility by Horizon",
        "Section 5 — Baseline behaviour",
        "Opportunity-scale diagnostic; Sharpe ratio is not applicable",
        "Future range and realized volatility rise nonlinearly with holding horizon, providing a clear unsigned opportunity target but no directional signal.",
        "The blue series shows mean future high-low range in ticks; the orange series shows mean realized volatility in basis points on a second axis. Both are calculated from complete forward paths. Mean range rises from about 28.7 ticks at 5 minutes to 50.7 at 15 minutes and 72.4 at 30 minutes, continuing upward at longer horizons. The two axes allow differently scaled measures to be compared by shape, not by absolute height.",
        "Read each series against its own y-axis. The common upward curvature reflects accumulation of movement, but neither metric should be expected to grow linearly with time because intraday volatility clusters and price increments partly offset. The fact that range and realized volatility co-move provides a useful cross-check: both respond to opportunity intensity through different definitions.",
        "This baseline motivates an expansion model. Forecasting large future range can improve timing, stop selection, or capital allocation even if expected signed return is zero. It also explains why ATR and realized-volatility features dominate later expansion IC rankings. Their strength is economically coherent, but it may be redundant: multiple volatility measures can encode the same latent state.",
        "Use the chart to choose realistic forecast horizons and to normalize economic thresholds. A horizon should offer enough movement to clear costs and stops without pushing label availability or holding time beyond the research contract. Compare session-specific and time-of-day charts before applying one global scale.",
        "Neither line is a payoff series. Conditional volatility and future range are unsigned risk/opportunity variables, so Sharpe ratio would have no meaningful numerator. The dual-axis display can exaggerate visual similarity; correlation, incremental IC, and out-of-sample validation are still required.",
    ),
    FigureBrief(
        "section05_baseline/05_session_movement.png",
        "Session Movement Comparison",
        "Section 5 — Baseline behaviour",
        "Session opportunity diagnostic; Sharpe ratio is not applicable",
        "New York offers consistently more future range than London at every horizon, so session is a first-order conditioning variable rather than a cosmetic label.",
        "The two lines show mean future range in ticks for entries in the London and New York execution windows. At 5 minutes the means are roughly 21 ticks for London and 33 for New York. By 60 minutes they are about 74 and 120 ticks; by 180 minutes about 133 and 200 ticks. Both curves rise with horizon, but the New York curve remains materially higher.",
        "Compare the vertical gap at the same horizon. It widens in tick terms as the horizon increases, showing that a pooled baseline would blend distinct opportunity scales. The curves describe average path range, not signed returns, so the higher New York line means more movement and risk—not automatically more profit.",
        "Session conditioning is necessary for thresholds, calibration, stops, and model evaluation. A single global expansion threshold would tend to label more New York observations as high opportunity and could understate unusual London moves. Similarly, a model trained on pooled rows could learn session identity as a shortcut rather than genuine within-session state.",
        "Use session-specific normalization and report Validation results per session. Where a pooled operational model is desired, include session explicitly and test whether performance remains after within-session centering. Risk rules should also reflect the larger New York movement scale while retaining the same causal decision-time inputs.",
        "Sharpe ratio is not appropriate because the curves are unsigned overlapping path statistics. They omit direction, entry policy, transaction costs, and sequential position constraints. Mean range is also tail-sensitive; medians and quantiles should accompany it when calibrating live risk.",
    ),
    FigureBrief(
        "section05_baseline/06_session_movement_and_availability.png",
        "Session Movement and Label Availability",
        "Section 5 — Baseline behaviour",
        "Opportunity and data-quality diagnostic; Sharpe ratio is not applicable",
        "New York has the larger movement opportunity, while both sessions retain near-complete labels; the economic session gap is much larger than the availability gap.",
        "The left panel compares mean future range by session and horizon. The right panel compares the availability rate of forward-return labels. New York ranges are consistently larger, reaching about 200 ticks at 180 minutes versus about 133 for London. Availability stays above roughly 98.5% for both sessions, with London declining slightly faster at intermediate long horizons.",
        "Read the panels together. If the range gap were accompanied by a large availability gap, it could be a data-selection artefact. Instead, label coverage remains extremely high, so the larger New York range is more plausibly a real session characteristic. The tight availability axis magnifies sub-percentage-point differences; interpret them using the tick labels rather than visual slope alone.",
        "The joint view strengthens the case for session-specific baselines and model cells. It also shows why missingness cannot explain the strong New York expansion relationships on its own. However, long-horizon London entries remain somewhat more exposed to incomplete paths, so exact horizon/session denominators and trading-date counts must stay attached to every estimate.",
        "Use this figure when setting horizon-specific quality gates. A model result should not advance if it is driven by a cell with materially worse path completion than its comparator. Operationally, session-specific thresholds and risk limits should be preferred over a global rule.",
        "No sequential payoff is represented, so Sharpe is not applicable. Near-100% availability does not eliminate overlap, serial dependence, or selection near the forced exit. The chart reports means and coverage, not tail risk or execution feasibility.",
    ),
    FigureBrief(
        "section05_baseline/07_time_of_day_range_heatmaps.png",
        "Time-of-Day Future-Range Heatmaps",
        "Section 5 — Baseline behaviour",
        "Clock-conditioned opportunity diagnostic; Sharpe ratio is not applicable",
        "Opportunity varies strongly within each session as well as across horizons, confirming that clock features encode real structure but may also be powerful baseline proxies.",
        "Each heatmap reports mean future range in ticks for an entry-time bin (rows) and horizon (columns). London occupies the left panel and New York the right. Darker cells indicate smaller expected range; yellow cells indicate larger range. Horizon drives a broad left-to-right increase, while vertical colour changes reveal intraday seasonality within a fixed horizon.",
        "Read across a row to see how movement accumulates from 5 to 180 minutes for one entry time. Read down a column to compare entry times at a fixed horizon. In New York, early-to-mid-morning entries generally show the largest long-horizon range, with later entries fading as less active time remains. London is more stable at short horizons, but the 180-minute column brightens near the end of the entry window because it reaches into more active later trading.",
        "The pattern explains why clock features can post extremely strong expansion IC: time remaining, session progress, and time-to-cutoff are structurally related to the future path. Such features are useful for calibration and opportunity forecasting, but they are not necessarily novel alpha. Redundancy analysis must separate genuine state information from deterministic clock exposure.",
        "Use time-of-day baselines for residualization, threshold setting, and monitoring. A live model should be evaluated within clock buckets so drift in market activity is not confused with feature improvement. When a clock feature dominates, compare against a transparent clock-only anchor before accepting a complex model.",
        "Sharpe ratio is inapplicable because the heatmap contains unsigned overlapping future ranges, not returns from a sequential strategy. The colour scales differ between panels, so colour intensity should not be compared directly across London and New York without consulting the bars.",
    ),
    FigureBrief(
        "section06_features/06_feature_family_inventory.png",
        "Engineered-Feature Family Inventory",
        "Section 6 — Feature engineering",
        "Research-design inventory; Sharpe ratio is not applicable",
        "The 85-feature library is broad but disciplined: 79 core predictors span conventional market-state families and only six are explicitly experimental.",
        "The horizontal bars count registered predictors by family. Price/return/momentum is the largest family with 16 features, followed by volatility/range state with 13 and trend/persistence with 11. Candle geometry, volume/activity, and VWAP/session state each contribute 10; clock/calendar contributes nine. The orange bar isolates the six experimental hypotheses from the total inventory.",
        "Read bar length as research breadth, not importance. A large family creates more opportunities for correlated variants and multiple-testing false positives. The registry includes causal availability timestamps, formula definitions, reset boundaries, data types, and tier labels, so the count reflects governed predictors rather than ad hoc notebook columns.",
        "The mix is economically coherent: conventional features establish transparent anchors, while a small experimental subset tests novel ideas without overwhelming the search space. The family counts also foreshadow redundancy. ATR ladders, realized-volatility windows, return horizons, and clock transforms are related by construction, so 85 columns do not represent 85 independent hypotheses.",
        "Use this inventory to enforce coverage, lineage, and family-aware validation. Shortlists should be compared with the original family mix to detect concentration. If one family dominates leaders, cluster and incremental-information tests are mandatory before multivariate modelling.",
        "Sharpe ratio cannot evaluate a raw feature inventory. A feature has no payoff until an orientation, signal threshold, holding rule, sizing rule, costs, and sequential portfolio logic are defined. Counting features also says nothing about predictive quality and increases the need for multiple-testing controls.",
    ),
    FigureBrief(
        "section06_features/07_feature_missingness.png",
        "Feature Missingness Diagnostics",
        "Section 6 — Feature engineering",
        "Data-quality diagnostic; Sharpe ratio is not applicable",
        "Feature coverage is uniformly strong: even the worst Development missingness is about 1.15%, with most of the top-20 list below one-half percent.",
        "The chart ranks the 20 highest Development missing rates among engineered predictors. Two-bar directional balance is the largest at 1.150%, followed by several session/VWAP features near 0.42–0.45%. The remaining displayed features are around 0.20% or lower. All reported non-null values passed the finite-value checks.",
        "Read the x-axis carefully: percentage labels are rounded, making several small bars appear as 0%. The ranking remains useful even when displayed ticks are coarse. Missingness mainly reflects causal lookback warm-up, session-state initialization, or unavailable denominators—not post-outcome filtering.",
        "Low missingness reduces the chance that model comparisons are driven by different row populations. Still, the affected features are not random: session-range and VWAP distances are missing early in a session, and directional-balance features require prior bars. Imputing them without a missingness indicator could blur exactly the boundary states they represent.",
        "Use Development-fitted imputation and preserve the same transformation in Validation and live inference. Monitor missingness by session and partition, not just globally. A feature should be quarantined if coverage changes materially, finite rates fall, or missingness correlates with the target through a data leak.",
        "Sharpe ratio is not a coverage measure. This figure precedes any payoff construction. Low missingness does not validate predictive value, and pairwise complete-case analysis can still create inconsistent samples across features if masks are not frozen.",
    ),
    FigureBrief(
        "section07_univariate/08_validation_ic_leaderboard.png",
        "Validation Information-Coefficient Leaderboard",
        "Section 7 — Univariate feature evaluation",
        "Primary predictive-association metric: daily rank IC",
        "Validation leaders are dominated by volatility and deterministic clock structure, with absolute daily rank ICs around 0.6–0.76; the sign encodes orientation, not quality.",
        "The bars rank the 20 strongest Validation feature/outcome/session/horizon cells by absolute mean daily Spearman rank IC. Red denotes negative IC and green positive IC. Expansion leaders include ATR and realized-volatility measures at 120–180 minutes; clock features such as minutes to noon or forced exit appear with opposite signs to equivalent progress measures. A New York 180-minute directional distance-from-session-open feature also ranks highly.",
        "Read magnitude as monotonic predictive association within daily groups. A negative bar can be just as useful as a positive one if the orientation is stable: for example, higher current ATR may imply lower ATR-normalized future expansion, producing a strong negative IC. Opposite-signed clock transforms can encode essentially the same information and should not be counted as independent discoveries.",
        "The chart confirms that expansion is forecastable out of sample, but it also highlights redundancy. ATR, realized volatility, session progress, and time remaining cluster around the same opportunity state. The notebook records 55 expanding-feature advancers and zero directional-feature advancers under the governed criteria, so this leaderboard supports opportunity forecasting rather than standalone directional alpha.",
        "Use IC to shortlist monotonic relationships, then require yearly stability, economic bucket spreads, cluster controls, partial IC beyond an anchor, and multivariate Validation improvement. Compare within the same session and horizon; ranking across unlike targets can mislead.",
        "Sharpe is not the primary metric here. IC evaluates association without requiring a trading policy. The feature-spread Sharpe in the next figure is a secondary diagnostic only because overlapping events, missing costs, and unconstrained positions prevent interpreting these cells as executable performance.",
    ),
    FigureBrief(
        "section07_univariate/09_feature_spread_sharpe.png",
        "Feature-Spread Sharpe Leaderboard",
        "Section 7 — Univariate feature evaluation",
        "Secondary screening Sharpe; explicitly not executable strategy Sharpe",
        "Several feature-sorted Validation spreads have extremely high diagnostic Sharpes, but the largest values often rest on few trading dates and must not be read as tradable performance.",
        "The chart ranks Validation cells by annualized Sharpe of a daily top-versus-bottom feature-bucket spread. Development determines the orientation; Validation supplies the confirmation view. The leading cell—distance from research-day VWAP, London, 180 minutes—has Sharpe 32.60 on only 33 dates. New York 180-minute VWAP slope reaches 30.33 on 96 dates. With at least 150 dates, session-range position (New York, 180 minutes) still reaches 23.74 and fraction above VWAP reaches 22.77.",
        "Read the bar length together with the hidden denominator: Sharpe equals mean daily spread divided by daily spread volatility, annualized by square-root 252. A high value can result from a large stable cross-sectional spread, but small date counts make the ratio fragile. The orientation low-minus-high is often dictated by negative directional IC.",
        "The diagnostic shows that some features consistently separate signed outcomes across extreme buckets. It does not prove a realizable strategy. Many observations overlap within a date; a top and bottom bucket can imply simultaneous or repeated positions; costs, capacity, stop rules, and position overlap are absent. The values are therefore useful for ranking robustness, not capital allocation.",
        "Use this chart to complement IC: prefer features with stable Validation Sharpe, adequate trading dates, consistent yearly signs, and incremental information. Apply a minimum-date filter and inspect sensitivity to bucket definitions. Advancement still depends on the governed IC/economic-spread rules.",
        "Do not compare these numbers with a conventional portfolio Sharpe. The metric status is `screening_diagnostic_not_executable_pnl`. Annualization assumes daily observations but cannot repair overlap inside a day or selection bias across 1,992 tested cells. Multiple-testing and winner's-curse risk are substantial.",
    ),
    FigureBrief(
        "section07_univariate/10_ic_vs_spread_sharpe.png",
        "Information Coefficient versus Spread Sharpe",
        "Section 7 — Univariate feature evaluation",
        "Joint predictive and economic-screening diagnostic",
        "IC and diagnostic Sharpe are related but far from interchangeable: association strength does not uniquely determine the stability or scale of an extreme-bucket spread.",
        "Each point is a Validation feature/session/horizon cell. The x-axis is absolute daily directional rank IC; the y-axis is annualized daily feature-spread Sharpe. London and New York use different markers. The cloud asks whether stronger monotonic rank association also produces a more stable daily separation between Development-oriented extreme buckets.",
        "Read points to the right as stronger rank ordering and points higher as more stable/economically larger daily spreads. A point high but left can arise when only the extremes separate while the full ranking is weak. A point right but lower can indicate a broad monotonic relationship with noisy or small extreme-bucket payoff. Session clusters reveal that the mapping differs by market context.",
        "The dispersion confirms that model evaluation needs more than one statistic. IC is less dependent on an arbitrary bucket cut and is the governed association metric. Spread Sharpe adds an intuitive risk-adjusted screen, but it is more sensitive to bucket occupancy, date count, tail outcomes, and annualization. Agreement between the two is stronger evidence than either alone; disagreement is a prompt for investigation.",
        "Use the plot to identify robust upper-right candidates and suspicious outliers. Annotate or tabulate candidates with minimum date counts, then check yearly stability, orientation, and partial IC. Avoid choosing a feature merely because it maximizes one axis.",
        "The y-axis remains a screening Sharpe, not strategy Sharpe. It excludes costs and sequential constraints. The x-axis uses absolute IC, so it suppresses sign; orientation must be recovered from the underlying cell before any signal is designed.",
    ),
    FigureBrief(
        "section07_univariate/11_yearly_ic_stability.png",
        "Year-by-Year Feature Stability",
        "Section 7 — Univariate feature evaluation",
        "Temporal robustness diagnostic based on annual daily rank IC",
        "Leading features that preserve sign and comparable magnitude across years are more credible than pooled winners driven by a single regime.",
        "The two panels plot annual mean daily directional rank IC for selected leading spread features at the 60-minute horizon, split into London and New York. Each line is a feature and each point a calendar year. The zero line separates orientation changes from stable relationships.",
        "Read consistency before height. A line that stays on one side of zero across Development years and the Validation year is more robust than a line with one extreme peak and sign reversals. New York clock features show large, persistent magnitudes around 0.43–0.54 in absolute value, while some London clock relationships hover near zero and change sign. Those differences reinforce session-specific interpretation.",
        "Temporal stability guards against pooled-sample illusion. Market volatility, session structure, and roll behaviour change over time; a feature that survives those shifts is likelier to encode structural state. Yet stable clock features can still be deterministic baseline effects rather than alpha, so redundancy and anchor comparisons remain necessary.",
        "Use this chart as a veto, not merely a ranking. A candidate that depends on one exceptional year should fail advancement even if pooled IC or spread Sharpe is impressive. Where the latest year weakens, investigate drift and recalibration rather than averaging it away.",
        "Sharpe is not shown because the object is annual predictive association, not a sequential return stream. With only a few annual points, visual stability is qualitative; formal date-block intervals and per-year trading-date counts should accompany decisions. Final-test years must not be used to retune selection.",
    ),
    FigureBrief(
        "section08_redundancy/12_redundancy_clusters.png",
        "Feature Redundancy Clusters",
        "Section 8 — Redundancy and incremental information",
        "Dependence-structure diagnostic; Sharpe ratio is not applicable",
        "The 85-feature registry collapses to 30 Development-fitted clusters, with the largest volatility cluster containing seven near-substitutable predictors.",
        "The bars show the 15 largest correlation clusters, labelled by cluster ID and representative feature. The largest cluster is represented by ATR(20) and contains seven ATR/realized-volatility variants. Other clusters group time-to-exit variables, execution-window clock transforms, range-compression measures, relative-volume measures, and VWAP-state features. Cluster membership is fitted on Development only.",
        "Read cluster size as redundancy, not predictive quality. A large bar means many registered features move similarly after daily aggregation. The representative is a convenient proxy, not necessarily the universal best feature. Small or singleton clusters can be genuinely distinct or simply weakly correlated; their value must be established by out-of-sample evidence.",
        "The chart explains why raw leaderboards overstate breadth. If several ATR and realized-volatility measures occupy top IC ranks, they may all reflect one latent volatility state. Feeding every variant into a model can inflate coefficient instability, multiple testing, and apparent feature importance without adding information.",
        "Use clusters to select transparent representatives, build family-aware models, and avoid double-counting evidence. When a non-representative feature appears stronger, test whether it adds partial IC beyond the anchor rather than swapping labels within the same cluster.",
        "Sharpe ratio cannot quantify redundancy. Correlation clusters describe predictors, not returns. Results depend on the Development sample, aggregation, distance threshold, and correlation measure; nonlinear redundancy may remain. Cluster structure should be monitored for drift before live deployment.",
    ),
    FigureBrief(
        "section08_redundancy/13_incremental_information.png",
        "Incremental Information beyond the Anchor",
        "Section 8 — Redundancy and incremental information",
        "Primary metric: partial daily rank IC after controlling for ATR(20)",
        "Only features that retain Validation partial IC after the ATR anchor provide evidence of genuinely incremental opportunity information.",
        "Each panel compares raw daily IC on the x-axis with partial daily IC on the y-axis after removing the anchor's information date by date. The left panel is Development and the right Validation. The diagonal marks equality; points below it lose association after conditioning. Green points pass the declared incremental gates. The analysis finds 14 confirmed incremental candidates among 29 candidates, with ATR(20) as anchor.",
        "Read vertical displacement from the diagonal as the amount of raw IC explained by the anchor. A point near the x-axis had an impressive univariate result that vanishes when volatility state is removed. Points retaining large same-sign partial IC are more distinct. Clock features and certain volatility ratios retain material Validation partial IC, while many absolute-volatility measures collapse.",
        "This is the bridge from broad screening to parsimonious modelling. It prevents a ridge model from receiving many aliases of the same state and makes improvement claims interpretable. Notably, no directional features advance under the governed screen, whereas 15 expansion features enter the frozen multivariate set.",
        "Use partial IC and its date-block uncertainty to freeze candidates before model fitting. Compare Development and Validation panels for sign retention and magnitude decay. A feature with strong raw IC but weak partial IC should remain a descriptive proxy, not be counted as a new source of information.",
        "Sharpe is not appropriate because the test evaluates incremental association for an unsigned expansion target. Partial IC is conditional and scale-free; it does not estimate a portfolio payoff. Linear/rank residualization may miss nonlinear interactions, so multivariate Validation improvement is still required.",
    ),
    FigureBrief(
        "section09_models/14_multivariate_regression.png",
        "Multivariate Regression Performance",
        "Section 9 — Multivariate research",
        "Primary metrics: Validation daily rank IC and paired IC improvement",
        "The ridge expansion model advances in three of four cells; London 180 minutes improves only 0.0082 IC and fails the predeclared 0.02 margin.",
        "The left panel compares Validation daily rank IC for the ATR anchor and the multivariate ridge model in London/New York at 60 and 180 minutes. The right panel plots paired model-minus-anchor IC improvement with date-block confidence intervals and a dashed 0.02 advancement margin. Improvements are 0.0453 for London 60m, 0.0082 for London 180m, 0.0810 for New York 60m, and 0.0876 for New York 180m.",
        "Read the left panel for absolute predictive strength and the right for incremental value. An interval above zero shows a statistically confirmed improvement; clearing 0.02 shows the gain is also large enough under the declared rule. London 180m is positive with interval [0.0035, 0.0130] but still anchor-sufficient because the point and interval remain below the materiality margin.",
        "The result supports a multivariate opportunity model in three cells while preventing a statistically significant but operationally small improvement from being oversold. New York benefits most, reaching model IC 0.626 at 60m and 0.849 at 180m. The high values reflect forecastability of relative expansion and clock/volatility structure, not direction.",
        "Use the advancing cells for downstream opportunity gates, with the selected ridge penalty and feature set frozen from Development walk-forward analysis. Retain the anchor for London 180m unless a new preregistered study shows material incremental value. Monitor both absolute IC and paired improvement after deployment.",
        "Sharpe is intentionally absent. The target is unsigned future range, and regression outputs are not positions. Converting model scores to trades requires direction, thresholds, costs, sizing, and a sequential backtest. High expansion IC can coexist with negative trading Sharpe, as the later backtest demonstrates.",
    ),
    FigureBrief(
        "section09_models/15_classification_performance.png",
        "Classification Performance",
        "Section 9 — Multivariate research",
        "Primary metrics: Validation AUC and Brier score",
        "The logistic model improves discrimination and probability error in every cell, but New York 60-minute upper-decile probabilities are overconfident and require recalibration before sizing use.",
        "The left panel compares model and anchor AUC; higher is better. The right compares Brier score; lower is better. Validation model AUC ranges from 0.800 to 0.928 and exceeds the anchor in all four cells. Model Brier scores also improve, including New York 180m at 0.084 versus 0.102 for the anchor.",
        "AUC measures ranking between expansion and non-expansion outcomes, while Brier score penalizes probability error. Similar AUC improvements can coexist with different calibration quality. The chart therefore combines discrimination with overall probability accuracy, but it does not expose decile-specific bias.",
        "The model is useful for opportunity classification, especially New York 180m. However, New York 60m overstates its upper probability bins: predictions around 0.65–0.82 correspond to realized rates around 0.41–0.55. Ranking may remain strong while probabilities are unsafe for direct risk scaling. London and New York 180m are more rank-consistent.",
        "Use AUC for ranking validation, Brier score and reliability curves for probability use, and recalibrate before converting probabilities into size. Thresholds must be fitted on Development and frozen. Report base rates because AUC can look strong even when positive classes are relatively uncommon.",
        "Sharpe ratio is inappropriate for an unsigned class probability. The chart evaluates forecasts, not realized returns. A good classifier can time movement but still lack direction or fail after costs. AUC is also insensitive to calibration, and Brier score is affected by class prevalence.",
    ),
    FigureBrief(
        "section10_signals/16_signal_diagnostics.png",
        "Signal Construction Diagnostics",
        "Section 10 — Statistical signal construction",
        "Pre-backtest rule diagnostic; Sharpe ratio becomes applicable only after realized sequential returns",
        "The frozen opportunity gate retains about one-fifth of candidates out of sample, while volatility-scaled stops frequently hit the 10-tick floor.",
        "The left panel compares Development and Validation gate rates by session. Development thresholds are the 80th percentile, producing about 20.0% retention by design; Validation rates are 21.86% in London and 20.98% in New York. The right panel shows stop distances from 1.5× ATR, clamped to 10–100 ticks, with a median of 10.575 ticks. Across 419,063 candidates, 85,362 pass the opportunity gate.",
        "Read the left panel as a distribution-shift check: close Development/Validation rates indicate the frozen score thresholds transfer reasonably. The right panel reveals the practical risk scale. Approximately 46% of stops are clamped, so the lower bound materially shapes the strategy rather than serving as a rare safety guard.",
        "This figure validates rule construction, not profitability. The model is an opportunity forecast; combining it with long and short benchmarks tests whether high predicted movement adds directional value. Similar gate rates cannot guarantee similar returns, and a heavily binding stop floor can alter payoff asymmetry and cost sensitivity.",
        "Use these diagnostics before every backtest and live run. Alert on gate-rate drift, stop-clamp frequency, and median risk distance. If thresholds or bounds change, treat it as a new strategy specification requiring fresh out-of-sample validation.",
        "Sharpe is premature here because realized sequential daily returns have not yet been generated. Candidate counts and risk distances are not payoffs. Once the one-position-at-a-time backtest applies entries, exits, costs, and daily aggregation, Sharpe becomes a valid primary metric.",
    ),
    FigureBrief(
        "section11_backtest/17_strategy_sharpe_by_cost.png",
        "Strategy Sharpe by Transaction-Cost Scenario",
        "Section 11 — Independent sequential backtest",
        "Primary valid Sharpe metric: annualized daily net-R Sharpe",
        "All base-cost and pessimistic variants have deeply negative Sharpe in both Development and Validation; transaction costs expose the absence of standalone directional edge.",
        "The panels compare four sequential strategy variants across frictionless, base, and pessimistic costs in Development and Validation. Returns are summed by New York trading date, including zero-trade dates, and annualized with 252 trading days. Under base costs, Development Sharpe ranges from -14.33 to -8.32 and Validation from -14.27 to -7.35. Both long/short and gated/ungated variants fail.",
        "Read each group from frictionless to pessimistic to see cost sensitivity. Moving below zero means the daily return stream loses money after risk adjustment; values this negative indicate persistent losses, not a marginal miss. The expansion gate reduces trade count but does not supply direction, so it cannot rescue the benchmark.",
        "This is the notebook's primary legitimate use of Sharpe because the backtest creates a chronological, one-position-at-a-time, cost-adjusted return stream. The agreement between Development and Validation is decisive: the standalone statistical system is rejected. The useful component is opportunity forecasting, which may assist an independently directional strategy.",
        "Do not tune toward the least-negative bar. Preserve the rejection and transfer only preregistered components to hybrid research. Any future strategy must demonstrate positive Validation Sharpe under base and pessimistic costs, with trade counts, drawdowns, and stability reported alongside it.",
        "Sharpe still has limitations: daily returns are non-normal, the sample contains 638 Development and 246 Validation dates, and annualization does not eliminate serial dependence. Costs are scenario assumptions, and Sharpe does not show tail losses or capital requirements. Profit factor, expectancy, and drawdown provide necessary complements.",
    ),
    FigureBrief(
        "section11_backtest/18_expectancy_vs_sharpe.png",
        "Expectancy versus Sharpe",
        "Section 11 — Independent sequential backtest",
        "Joint trade-level and daily time-series performance diagnostic",
        "Every base-cost variant lies in the negative-expectancy, negative-Sharpe quadrant, so trade economics and daily risk adjustment deliver the same rejection.",
        "Each point represents a base-cost variant/partition combination. The x-axis is mean net R per trade; the y-axis is annualized daily net-R Sharpe. Development and Validation use different markers. Mean trade expectancy ranges roughly from -0.19R to -0.28R, while daily Sharpe ranges from about -7.35 to -14.33.",
        "The zero lines create four quadrants. A credible strategy should appear in the upper-right: positive average trade edge and positive risk-adjusted daily performance. All points are lower-left. Gated strategies reduce the number of trades and may have less negative daily Sharpe, but their mean trade expectancy is often worse, showing that filtering on expansion does not create direction.",
        "Agreement across frequencies is important. Expectancy measures average trade economics; Sharpe measures how the sequence aggregates by day relative to daily volatility. Divergence could reveal clustering or a few high-volume days, but here both metrics reject the strategy in both partitions. That makes the conclusion robust to the performance lens.",
        "Use this chart to avoid optimizing one metric in isolation. Require positive expectancy after costs, positive daily Sharpe, reasonable drawdown, and stable yearly results. A point moving upward only because trading becomes sparse should be checked for active-day versus calendar-day treatment.",
        "Both metrics inherit the backtest assumptions. Expectancy ignores ordering and tail clustering; Sharpe compresses the path into mean/volatility. Neither captures capacity or execution slippage beyond scenarios. The plotted annotations can overlap, so the underlying table remains authoritative.",
    ),
    FigureBrief(
        "section11_backtest/19_daily_equity_curves.png",
        "Daily Equity Curves",
        "Section 11 — Independent sequential backtest",
        "Path diagnostic at the same daily frequency used for Sharpe",
        "All base-cost equity curves decline persistently in both partitions, confirming that the negative Sharpe is produced by broad pathwise losses rather than one isolated event.",
        "The two panels plot cumulative net R by New York trading date for the four base-cost variants in Development and Validation. Daily aggregation matches the Sharpe calculation and includes the full candidate calendar. Ungated variants trade more and therefore accumulate larger absolute losses; gated variants fall more slowly but remain decisively negative.",
        "Read slope, drawdown episodes, and recovery behaviour. A steadily negative slope means repeated adverse expectancy. Sharp drops indicate clustered loss days; flat periods can reflect fewer trades rather than stability. Compare Development and Validation shapes rather than only endpoints: similar deterioration suggests the failure generalizes out of sample.",
        "The paths support the conclusion from Sharpe and expectancy. There is no hidden profitable subperiod large enough to challenge rejection, and the gate acts mostly as trade throttling. Because the curves are in R units, they express loss relative to per-trade risk, but not a deployable account balance or margin model.",
        "Use equity curves to diagnose regime dependence, operational exposure, and recovery time. Overlay major rule changes only in a new governed study. If a future variant has positive Sharpe but an unacceptable path, drawdown and concentration should veto deployment.",
        "Cumulative curves can exaggerate differences in trade count and do not normalize by capital or active risk. They also omit confidence bands. Sharpe summarizes the same daily increments but cannot replace visual path inspection; both must be interpreted with costs, trade counts, and drawdown.",
    ),
    FigureBrief(
        "section11_backtest/20_profit_factor_drawdown.png",
        "Profit Factor and Maximum Drawdown",
        "Section 11 — Independent sequential backtest",
        "Complementary payoff-efficiency and path-risk metrics",
        "Every base-cost profit factor is below one and every variant experiences severe drawdown, corroborating the strongly negative daily Sharpe.",
        "The left panel plots gross gains divided by gross losses for each variant/partition. The dashed line at one is break-even before considering capital path: values below one mean losses exceed gains. Base-cost profit factors range roughly from 0.67 to 0.76. The right panel plots maximum trade-sequence drawdown in R, with ungated Development variants experiencing losses of several thousand R because of high trade count and negative expectancy.",
        "Read profit factor as payoff efficiency and drawdown as path severity. A strategy can have positive Sharpe with an unattractive drawdown or a profit factor near one with unstable timing; here both are unambiguously poor. Gating reduces total exposure and therefore absolute drawdown, but it does not improve the underlying edge enough to cross profit factor one.",
        "These metrics explain the negative Sharpe economically. Win rates around one third at a 2R target resemble directionless behaviour; costs lower net wins and deepen losses. The strategy therefore compounds losses as trading frequency increases. Lower drawdown for a gated variant should not be mistaken for success when the curve still trends down.",
        "Use a multi-metric acceptance rule: positive base-cost expectancy, profit factor above one with a safety margin, positive Validation Sharpe, and drawdown within capital limits. Normalize drawdown by intended risk budget before deployment; raw R totals mainly compare variants under the shared one-R convention.",
        "Profit factor ignores time and can be dominated by a few large trades. Maximum drawdown depends on start date, ordering, and trade frequency. Neither replaces Sharpe, and Sharpe does not replace them. All remain simulated results subject to execution and cost assumptions.",
    ),
    FigureBrief(
        "section12_hybrid/21_hybrid_improvements.png",
        "Hybrid Integration Improvements",
        "Section 12 — Hybrid integration research",
        "Date-block event-study effect; Sharpe ratio is not applicable",
        "Continuation-short appears positive in both partitions, but the gated sample is too small and its Validation interval crosses zero, so every hybrid family is classified as no incremental value.",
        "The plot shows gated-minus-baseline mean robust R for continuation/reversal and long/short families in Development and Validation, with date-block confidence intervals. Continuation-short improves by +1.17R in Development and +0.87R in Validation, but only 410 and 241 gated events across 33 and 30 dates support those estimates. Other families are negative or unstable.",
        "Read the point as effect size and the horizontal interval as date-level uncertainty. Crossing zero means the improvement is not confirmed. Even a point with an interval above zero can fail preregistered event and date floors. The y-axis labels combine hypothesis, side, and partition so sign retention can be assessed directly.",
        "The chart is an instructive small-sample warning. The most attractive mean improvement survives in sign but not in governed evidence. The opportunity gate passes only about 5% of POI events because those events already form in extended conditions; conditioning creates a thin, selected subset. The notebook therefore refuses to convert a dramatic point estimate into a trading claim.",
        "Use the result to close the tested hybrid families and preserve the sample-floor veto. Future work should preregister a different interaction or acquire more independent dates rather than retuning the existing gate. Report event counts and trading dates next to every effect.",
        "Sharpe is inappropriate because events overlap and the robust-R difference is not a sequential portfolio return stream. Annualizing these effects would create false precision. The correct metrics are effect size, date-block interval, retention, and sample floors.",
    ),
    FigureBrief(
        "section12b_conditioning/22_sizing_effects.png",
        "Opportunity-Conditioned Sizing Effects",
        "Section 12B — Sizing, exits, and suppression",
        "Conditional event-study ratio effect; Sharpe ratio is not applicable",
        "Proportional sizing shows stable positive mean/MAD effects, but confidence and materiality rules fail; no sizing scheme advances.",
        "The left panel shows Development change in the mean-to-MAD ratio for proportional and inverse sizing under frictionless, base, and pessimistic costs, with date-block intervals. The right plots Development versus Validation effects; the diagonal represents full retention and colour indicates the criteria verdict. Proportional effects are +0.026 to +0.038 in Development and +0.028 to +0.031 in Validation, with retention 0.81–1.08.",
        "Read the left interval against zero and the right point against both axes and the diagonal. Proportional sizing is directionally consistent, while inverse sizing is consistently negative. However, the base Development interval includes zero and the mean-R shift is only about 0.04–0.09R, below the declared materiality requirement. Hence red verdicts despite good retention.",
        "This is the strongest persistent signal produced by the opportunity model, but the governance rule prevents relabelling a small concentration effect as a sizing breakthrough. The effect comes from overweighting high predicted-expansion quintiles, which is related to the thin event-study evidence already rejected in Section 12.",
        "Use the finding as a research clue, not a live sizing rule. Preserve equal-risk sizing in the approved system. A future test would need independent data, a preregistered materiality bound, and a true sequential account-level simulation with turnover and leverage constraints.",
        "Sharpe is not valid for these overlapping conditional event outcomes. Mean/MAD is a robust ratio within the event study, not a time-series portfolio statistic. Stable retention does not overcome an interval crossing zero or an economically immaterial mean shift.",
    ),
    FigureBrief(
        "section12b_conditioning/23_exit_and_suppression.png",
        "Exit-Horizon and Suppression Effects",
        "Section 12B — Sizing, exits, and suppression",
        "Conditional event-study effects; Sharpe ratio is not applicable",
        "Both adaptive exits and bottom-quintile suppression improve Development but reverse in Validation, providing a clear out-of-sample veto.",
        "The left panel shows H2 conditional-minus-unconditional mean R for adaptive exit horizons. Development is +0.0225R, while Validation is -0.0444R, a retention near -2.0. The right panel shows H3 suppression effects: Development median improvement +0.0607R and stop-rate reduction +0.0059, versus Validation -0.0048R and -0.0028.",
        "Read sign consistency first. Positive Development bars followed by negative Validation bars indicate the mapping did not generalize. The H2 unconditional best horizon is 120 minutes; a Development-fitted conditional menu selecting 60/120/240 minutes adds complexity without durable benefit. For H3, both co-primary effects must improve, so either Validation reversal is sufficient to reject.",
        "The figure demonstrates why holdout validation must control research enthusiasm. Effects that look plausible in-sample—shorter exits for low opportunity and suppressing the lowest quintile—can exploit noise or regime-specific paths. Because both hypotheses reverse, more tuning on the same Validation data would compound overfitting.",
        "Use the result to retain the unconditional 120-minute exit and avoid the tested suppression rule. Close the hypothesis family under the current protocol. Any revisit should use a new frozen sample and a simplified mechanism justified independently of these failed mappings.",
        "Sharpe is inapplicable because the bars are non-sequential conditional event effects. They do not include a portfolio calendar, overlapping-position resolution, capital, or complete costs. Effect sizes are small, and H2 Development uncertainty includes zero, further weakening the case.",
    ),
    FigureBrief(
        "garch_volatility/24_garch_regimes.png",
        "GARCH Volatility Regime History",
        "GARCH volatility modelling",
        "Risk-state forecast diagnostic; Sharpe ratio is not applicable",
        "The Development-frozen GARCH threshold identifies persistent volatility regimes and a marked rise in high-volatility exposure in the later sample.",
        "The upper panel plots daily median GC GARCH(1,1) conditional volatility with the Development 75th-percentile threshold. The lower panel shows the daily share of tradable bars classified as high volatility. The GC fit has alpha 0.1121, beta 0.8802, and persistence 0.9924, indicating slow decay of volatility shocks. The threshold is approximately 0.000250.",
        "Read excursions above the dashed line as risk-state episodes, not price direction. The lower panel converts the minute-level regime into daily prevalence: a day can be partly high-vol rather than receiving a single label. Clusters and long elevated periods are expected from the high estimated persistence.",
        "The regime feature captures time-varying risk that can support normalization, stops, exposure limits, or model monitoring. GC's high-vol rate rises from 23.4% in Development to 25.0% in Validation and 54.2% in Final test, while mean conditional volatility increases from 0.000214 to 0.000335 by Final test. That is material distribution shift and should trigger careful calibration review.",
        "Use the frozen threshold as a monitoring reference, not a target to refit opportunistically. Compare model errors, gate rates, costs, and drawdowns by regime. If live high-vol prevalence stays structurally above Development, reassess risk limits and normalization under a new governance cycle.",
        "Conditional volatility is an unsigned forecast, so Sharpe ratio has no meaningful numerator. A high-vol regime may contain profitable or losing moves in either direction. GARCH assumptions, parameter stability, and the MGC convergence warning also limit interpretation; regime labels are model-dependent.",
    ),
    FigureBrief(
        "garch_volatility/25_garch_intraday_regimes.png",
        "Intraday GARCH Returns and Regimes",
        "GARCH volatility modelling",
        "Risk-state visualization; Sharpe ratio is not applicable",
        "High-volatility classifications align with clusters of large 15-minute returns and elevated conditional variance, making the regime economically interpretable but not directional.",
        "The upper panel shows GC 15-minute log returns, coloured red when the underlying one-minute GARCH feature is in the high-volatility regime. The lower panel shows conditional volatility with the frozen Development p75 threshold and shaded high-vol periods. The data are resampled for visualization; the model itself is fitted to the one-minute series.",
        "Read the panels vertically. Red return clusters should coincide with orange volatility above the threshold below. GARCH volatility reacts to large shocks and decays slowly because estimated persistence is near one. The model therefore groups turbulent periods rather than treating each large bar as an isolated event.",
        "The alignment provides face validity for the feature. It can condition risk and opportunity models even when signed returns average near zero. The later sample contains more sustained high-volatility periods, consistent with the partition summaries. That shift can change stop distributions, cost impact, and score calibration even if the predictor remains causally available.",
        "Use this chart for regime diagnostics, anomaly review, and communication with risk stakeholders. Evaluate forecast errors and strategy metrics separately inside normal and high-vol states. Do not infer that red periods should be traded long or short without a validated directional mechanism.",
        "Sharpe is not applicable because neither conditional volatility nor regime labels are returns. The plot is dense and resampled, so individual bars should not be used for event attribution. GARCH is a parametric model; structural breaks, leverage effects, and non-Gaussian tails may require alternative specifications.",
    ),
)


def _set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths_dxa: Iterable[int]) -> None:
    widths = list(widths_dxa)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            _set_cell_width(cell, width)


def _set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_font(run, *, size: float | None = None, color: RGBColor | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _paragraph_bottom_rule(paragraph, color: str = RULE_GRAY, size: str = "8") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def _add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    _set_font(run, size=9, color=GRAY)
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def _configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = NAVY
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = GRAY
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _configure_section(doc: Document, brief: FigureBrief) -> None:
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run("STATISTICAL FEATURE RESEARCH  |  FIGURE INTERPRETATION")
    _set_font(run, size=8.5, color=GRAY, bold=True)
    _paragraph_bottom_rule(p)

    footer = section.footer
    p = footer.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    _add_page_field(p)


def _add_title_block(doc: Document, brief: FigureBrief) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(brief.title)
    _set_font(run, size=23, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(brief.section)
    _set_font(run, size=12.5, color=GRAY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    label = p.add_run("METRIC ROLE  ")
    _set_font(label, size=9, color=BLUE, bold=True)
    value = p.add_run(brief.metric_role)
    _set_font(value, size=9, color=GRAY)

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.14)
    p.paragraph_format.right_indent = Inches(0.14)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.08
    _paragraph_shading(p, LIGHT_GRAY)
    lead = p.add_run("Primary takeaway. ")
    _set_font(lead, size=10.5, color=BLUE, bold=True)
    body = p.add_run(brief.takeaway)
    _set_font(body, size=10.5, color=NAVY)


def _image_dimensions(path: Path, max_width: float = 6.35, max_height: float = 4.75) -> tuple[float, float]:
    with Image.open(path) as image:
        width_px, height_px = image.size
    aspect = width_px / height_px
    width = max_width
    height = width / aspect
    if height > max_height:
        height = max_height
        width = height * aspect
    return width, height


def _add_figure(doc: Document, brief: FigureBrief, image_path: Path) -> None:
    width, height = _image_dimensions(image_path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    inline_shape = run.add_picture(str(image_path), width=Inches(width), height=Inches(height))
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", brief.title)
    doc_pr.set("descr", f"{brief.title}. {brief.takeaway}")

    caption = doc.add_paragraph(style="Caption")
    caption.add_run(f"Figure source: {brief.image.replace('/', ' / ')}")


def _add_analysis(doc: Document, brief: FigureBrief) -> None:
    doc.add_page_break()
    sections = (
        ("What the figure represents", brief.represents),
        ("How to read it", brief.reading),
        ("Quantitative interpretation", brief.interpretation),
        ("Research and model implications", brief.implications),
        ("Limitations and correct metric use", brief.limitations),
    )
    for heading, text in sections:
        doc.add_heading(heading, level=2)
        p = doc.add_paragraph(text)
        p.paragraph_format.widow_control = True


def _build_document(brief: FigureBrief) -> Path:
    image_path = FIGURE_ROOT / brief.image
    if not image_path.is_file():
        raise FileNotFoundError(image_path)
    doc = Document()
    _configure_styles(doc)
    _configure_section(doc, brief)
    _add_title_block(doc, brief)
    _add_figure(doc, brief, image_path)
    _add_analysis(doc, brief)

    output_path = FIGURE_ROOT / f"{image_path.stem}_interpretation.docx"
    core = doc.core_properties
    core.title = f"{brief.title} — Quantitative Figure Interpretation"
    core.subject = "Statistical feature research figure interpretation"
    core.author = "Quant Research"
    core.keywords = "quantitative research, feature engineering, validation, Sharpe ratio"
    core.comments = "Generated from the executed statistical_feature_research notebook figure pack."
    doc.save(output_path)
    return output_path


def main() -> None:
    pngs = sorted(path.relative_to(FIGURE_ROOT).as_posix() for path in FIGURE_ROOT.rglob("*.png"))
    configured = sorted(brief.image for brief in BRIEFS)
    if pngs != configured:
        missing = sorted(set(pngs) - set(configured))
        stale = sorted(set(configured) - set(pngs))
        raise RuntimeError(f"Figure mapping mismatch. Unmapped={missing}; missing_source={stale}")

    records: list[dict[str, str]] = []
    for brief in BRIEFS:
        output_path = _build_document(brief)
        records.append(
            {
                "source_image": brief.image,
                "document": output_path.relative_to(FIGURE_ROOT).as_posix(),
                "title": brief.title,
                "metric_role": brief.metric_role,
            }
        )
        print(f"Created {output_path.relative_to(PROJECT_ROOT)}")

    manifest_path = FIGURE_ROOT / "figure_document_manifest.csv"
    with manifest_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=records[0].keys())
        writer.writeheader()
        writer.writerows(records)
    print(f"Created {manifest_path.relative_to(PROJECT_ROOT)}")
    print(f"Documents created: {len(records)}")


if __name__ == "__main__":
    main()
